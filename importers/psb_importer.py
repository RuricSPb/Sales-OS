from openpyxl import load_workbook

from models.object import Object
from models.company import Company


class PSBImporter:

    COLUMNS = {

        "psb_link": "Ссылка",
        "psb_id": "Id номер",

        "name": "Объект",
        "description": "Описание",

        "region": "Регион",
        "district": "Район",
        "address": "Адрес",

        "object_type": "Тип объекта",
        "work_type": "Вид работ",

        "stage": "Этап строительства",
        "visit_date": "Дата посещения объекта",

        "site": "Сайт",

        "construction_period": "Сроки строительства",

        "area_site": "Участок",
        "building_area": "Площадь застройки",
        "total_area": "Общая площадь",
        "floors": "Этажность",

        "frame_type": "Каркас",
        "foundation": "Фундамент",
        "walls": "Стены",
        "windows": "Окна",
        "roof": "Кровля",
        "floors": "Полы",
        "doors": "Двери",

        "company_name": "Название компании-участника строительства",
        "company_type": "Тип компании",
        "contact": "Контакт компании",
        "phone": "Тел. компании",
        "email": "E-mail компании",
        "inn": "ИНН компании"

    }


    def import_objects(self, filename):

        workbook = load_workbook(filename)

        sheet = workbook.active


        headers = {}

        for col in range(1, sheet.max_column + 1):

            value = sheet.cell(1, col).value

            if value:

                headers[str(value).strip()] = col


        objects = {}


        for row in range(2, sheet.max_row + 1):

            psb_id = self._value(
                sheet,
                row,
                headers,
                "psb_id"
            )


            if not psb_id:
                continue


            if psb_id not in objects:


                obj = Object(

                    name=self._value(
                        sheet,
                        row,
                        headers,
                        "name"
                    ),

                    source="ПСБ",

                    external_id=psb_id,

                    psb_id=psb_id,

                    psb_link=self._value(
                        sheet,
                        row,
                        headers,
                        "psb_link"
                    ),

                    description=self._value(
                        sheet,
                        row,
                        headers,
                        "description"
                    ),


                    region=self._value(
                        sheet,
                        row,
                        headers,
                        "region"
                    ),

                    district=self._value(
                        sheet,
                        row,
                        headers,
                        "district"
                    ),

                    address=self._value(
                        sheet,
                        row,
                        headers,
                        "address"
                    ),


                    object_type=self._value(
                        sheet,
                        row,
                        headers,
                        "object_type"
                    ),

                    work_type=self._value(
                        sheet,
                        row,
                        headers,
                        "work_type"
                    ),


                    stage=self._value(
                        sheet,
                        row,
                        headers,
                        "stage"
                    ),

                    visit_date=self._value(
                        sheet,
                        row,
                        headers,
                        "visit_date"
                    ),


                    frame_type=self._value(
                        sheet,
                        row,
                        headers,
                        "frame_type"
                    ),

                    construction_period=self._value(
                        sheet,
                        row,
                        headers,
                        "construction_period"
                    )

                )


                objects[psb_id] = {

                    "object": obj,

                    "companies": []

                }



            company_name = self._value(
                sheet,
                row,
                headers,
                "company_name"
            )


            if company_name:


                company = Company(

                    company_name=company_name,

                    company_type=self._value(
                        sheet,
                        row,
                        headers,
                        "company_type"
                    ),

                    contact=self._value(
                        sheet,
                        row,
                        headers,
                        "contact"
                    ),

                    phone=self._value(
                        sheet,
                        row,
                        headers,
                        "phone"
                    ),

                    email=self._value(
                        sheet,
                        row,
                        headers,
                        "email"
                    )

                )


                # временно сохраняем ИНН
                company.inn = self._value(
                    sheet,
                    row,
                    headers,
                    "inn"
                )


                objects[psb_id]["companies"].append(company)



        return objects



    def _value(
        self,
        sheet,
        row,
        headers,
        field
    ):

        header = self.COLUMNS.get(field)


        if header is None:

            return ""


        column = headers.get(header)


        if column is None:

            return ""


        value = sheet.cell(
            row=row,
            column=column
        ).value


        if value is None:

            return ""


        return str(value).strip()